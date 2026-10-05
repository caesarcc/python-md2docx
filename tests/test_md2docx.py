import tempfile
from pathlib import Path
import docx
import pytest

from md2docx.cli import parse_args, main
from md2docx.core import (
    find_input_files,
    resolve_output_path,
    clean_inline_math,
    build_docx,
    SUPPORTED_EXTENSIONS,
)


def test_find_input_files_single_file(tmp_path):
    f = tmp_path / "sample.md"
    f.write_text("# Hello", encoding="utf-8")
    assert find_input_files(f) == [f]


def test_find_input_files_directory(tmp_path):
    f1 = tmp_path / "a.md"
    f2 = tmp_path / "b.markdown"
    f3 = tmp_path / "c.txt"
    f1.write_text("# A", encoding="utf-8")
    f2.write_text("# B", encoding="utf-8")
    f3.write_text("C", encoding="utf-8")

    files = find_input_files(tmp_path, pattern="*.*")
    assert f1 in files
    assert f2 in files
    assert f3 not in files


def test_find_input_files_not_found(tmp_path):
    with pytest.raises(FileNotFoundError):
        find_input_files(tmp_path / "non_existent")


def test_resolve_output_path_overwrite(tmp_path):
    src = tmp_path / "doc.md"
    src.write_text("# Hi", encoding="utf-8")
    out = resolve_output_path(src, overwrite=True)
    assert out == tmp_path / "doc.docx"


def test_resolve_output_path_increment(tmp_path):
    src = tmp_path / "doc.md"
    src.write_text("# Hi", encoding="utf-8")

    f1 = tmp_path / "doc.docx"
    f1.write_text("dummy", encoding="utf-8")

    out2 = resolve_output_path(src, overwrite=False)
    assert out2 == tmp_path / "doc (2).docx"

    f2 = tmp_path / "doc (2).docx"
    f2.write_text("dummy", encoding="utf-8")

    out3 = resolve_output_path(src, overwrite=False)
    assert out3 == tmp_path / "doc (3).docx"


def test_clean_inline_math():
    raw = r"$\text{Ativos} \cap \text{Quitados} = \emptyset$ and $A \times B \to C$"
    cleaned = clean_inline_math(raw)
    assert "∩" in cleaned
    assert "∅" in cleaned
    assert "×" in cleaned
    assert "→" in cleaned
    assert r"\text" not in cleaned


def test_build_docx_complete(tmp_path):
    md_content = """# Título Principal

Texto de introdução com **negrito**, *itálico*, `código inline` e [link](https://example.com).

---

## 1. Seção de Teste

> Este é um blockquote estilizado.

- Item 1
  - Subitem 1.1
- [x] Tarefa concluída
- [ ] Tarefa pendente

| Coluna 1 | Coluna 2 |
| :--- | :---: |
| Dado A | Dado B |

```python
def hello():
    return "world"
```
"""
    src_file = tmp_path / "test_doc.md"
    src_file.write_text(md_content, encoding="utf-8")

    dest_file = tmp_path / "test_doc.docx"
    actual_dest = build_docx(
        source=src_file,
        dest=dest_file,
        orientation="vertical",
        font_size=10,
        page_size="a4",
        overwrite=True,
    )

    assert actual_dest.exists()
    doc = docx.Document(str(actual_dest))
    assert len(doc.paragraphs) > 0
    assert len(doc.tables) >= 2  # Table for blockquote/code and markdown table


def test_build_docx_with_diagram(tmp_path):
    md_content = """# Documento com Diagrama

```mermaid
pie title Distribuição de Teste
    "Item A" : 50
    "Item B" : 50
```
"""
    src_file = tmp_path / "diagram_doc.md"
    src_file.write_text(md_content, encoding="utf-8")

    dest_file = tmp_path / "diagram_doc.docx"
    actual_dest = build_docx(
        source=src_file,
        dest=dest_file,
        overwrite=True,
    )

    assert actual_dest.exists()
    # Check that an SVG file was created in output directory
    svg_files = list(tmp_path.glob("*.svg"))
    assert len(svg_files) >= 1
    assert svg_files[0].stat().st_size > 0


def test_cli_parse_args():
    args = parse_args(["docs", "-o", "horizontal", "-f", "12", "--page-size", "legal", "-w"])
    assert args.input == "docs"
    assert args.orientation == "horizontal"
    assert args.font_size == 12
    assert args.page_size == "legal"
    assert args.overwrite is True
